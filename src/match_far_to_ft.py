"""
match_far_to_ft.py

Resolves the join-key problem between the confirmed root-cause log (FAR)
and the bulk Final Tester (FT) export data.

Background
----------
The FAR log (confirmed, teardown-verified root causes) and the bulk FT
export do not share a common per-unit identifier: none of the FT export's
ID fields (unique_number, id_unique_number, lot_major, lot_minor,
device_id1/2/3, sc_ident) appear anywhere in the FAR log, and vice versa.

The fix: each FAR row carries a Purchase Order (PO) number. Contrinex can
pull a small FT export slice for a given PO — these PO-named files
(e.g. "1566194_CTX_LK_inductive_05.xlsx") contain the lot of units tested
under that PO, including their unique_number / id_unique_number / lot_major
fields. So PO becomes the join key between "a confirmed defect happened in
this PO" and "here are the FT test records for that PO's lot".

What this script does
----------------------
1. Reads the FAR master list (FAR_new_list.xlsx) and normalises rows.
2. Filters out rows that aren't usable for training:
     - "no deviation found" rows (teardown found nothing wrong -- not a
       root cause label)
     - rows whose article number isn't among the FT export files already
       collected
     - rows whose defect is not test-data-inferable in principle (e.g. PO
       2004884: repeated operator-handling issue -- confirmed by checking
       that lot's FT data and finding no electrical signal for it at all)
3. Cross-references each remaining row's PO against the PO-named FT files
   on disk and flags whether that PO's FT data is already available.
4. Writes data/processed/usable_far_rows.csv -- the shortlist that
   src/parser.py / src/features.py will build the labeled training table
   from once the flagged units are extracted per lot.

Usage
-----
    python src/match_far_to_ft.py \
        --far-list data/raw/FAR_new_list.xlsx \
        --ft-dir data/raw/ft_exports \
        --po-dir data/raw/po_reports \
        --out data/processed/usable_far_rows.csv

Both --ft-dir (bulk FT export files, named like "330-020-355...xlsx") and
--po-dir (PO-linked FT slices, named like "1566194_CTX_LK_inductive_05.xlsx")
are scanned for the article numbers / PO numbers they contain.
"""

import argparse
import csv
import glob
import os
import re
from collections import Counter

import openpyxl

# PO numbers to exclude outright: confirmed, via inspection of their FT
# data, to be process/operator-driven defects with no electrical signature
# -- not something a test-data classifier could ever learn.
EXCLUDED_POS = {
    "2004884",  # repeated unskilled-operator flex-damage lot; verified
                # against its 110-unit FT export: every unit passes
                # electrically (no ILedOn / isolation / eeprom flags),
                # so there is nothing here for a model to learn from.
}

NO_DEVIATION_RE = re.compile(r"no\s*deviation", re.IGNORECASE)


def article_number_from_filename(path: str) -> str | None:
    """Extract a normalised NNN-NNN-NNN article number from an FT filename."""
    base = os.path.basename(path)
    m = re.match(r"^(\d{3})[-_ ]?(\d{3})[-_ ]?(\d{3})", base)
    if not m:
        return None
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"


def po_from_filename(path: str) -> str:
    """PO-linked FT files are named '<PO>_CTX_LK_inductive_NN.xlsx'."""
    return os.path.basename(path).split("_")[0]


def collect_ft_articles(ft_dir: str) -> set[str]:
    articles = set()
    for f in glob.glob(os.path.join(ft_dir, "**", "*.xlsx"), recursive=True):
        art = article_number_from_filename(f)
        if art:
            articles.add(art)
    return articles


def collect_available_pos(po_dir: str) -> set[str]:
    pos = set()
    for f in glob.glob(os.path.join(po_dir, "**", "*.xlsx"), recursive=True):
        pos.add(po_from_filename(f))
    return pos - EXCLUDED_POS


def read_far_rows(far_list_path: str) -> list[list]:
    """
    Read the FAR master list. Handles a quirk in the export: the first
    data row is missing a leading blank column that every subsequent row
    has (an artifact of how the source report was generated), so it's
    detected and padded to line up with the rest.
    """
    wb = openpyxl.load_workbook(far_list_path, data_only=True, read_only=True)
    ws = wb.active

    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        row = list(row)
        if row and isinstance(row[0], str) and row[0].startswith("Select Item"):
            row = [None] + row[:-1]
        rows.append(row)

    # Column layout (0-indexed), once aligned:
    #  1: item id  2: date  3: CC reference  4: analyst  5: status
    #  6: PO       7: article no   8/9: designation
    #  10: series  11: symptom     12: root cause category
    #  13: root cause narrative    14: qty
    return [r for r in rows if len(r) > 3 and isinstance(r[3], str) and r[3].startswith("CC-")]


def build_usable_rows(far_rows: list[list], ft_articles: set[str], available_pos: set[str]):
    usable = []
    dropped_no_deviation = 0
    dropped_no_ft_article = 0
    dropped_excluded_po = 0

    for r in far_rows:
        cc = r[3]
        po = str(r[6]).strip() if r[6] else None
        article = str(r[7]).strip() if r[7] else None
        cause_category = r[12]
        narrative = r[13]
        qty = r[14] if len(r) > 14 else None

        if cause_category and NO_DEVIATION_RE.search(str(cause_category)):
            dropped_no_deviation += 1
            continue

        if not article or article not in ft_articles:
            dropped_no_ft_article += 1
            continue

        if po in EXCLUDED_POS:
            dropped_excluded_po += 1
            continue

        usable.append(
            {
                "CC_reference": cc,
                "PO": po,
                "article_no": article,
                "series": r[10],
                "symptom": r[11],
                "root_cause_category": cause_category,
                "root_cause_narrative": narrative,
                "qty": qty,
                "po_already_covered": po in available_pos,
            }
        )

    stats = {
        "total_far_rows": len(far_rows),
        "dropped_no_deviation": dropped_no_deviation,
        "dropped_no_ft_article": dropped_no_ft_article,
        "dropped_excluded_po": dropped_excluded_po,
        "usable_rows": len(usable),
        "po_covered": sum(1 for u in usable if u["po_already_covered"]),
    }
    return usable, stats


def write_csv(rows: list[dict], out_path: str):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    fieldnames = [
        "CC_reference", "PO", "article_no", "series", "symptom",
        "root_cause_category", "root_cause_narrative", "qty",
        "po_already_covered",
    ]
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--far-list", required=True, help="Path to FAR_new_list.xlsx")
    parser.add_argument("--ft-dir", required=True, help="Directory of bulk FT export files")
    parser.add_argument("--po-dir", required=True, help="Directory of PO-linked FT files")
    parser.add_argument("--out", default="data/processed/usable_far_rows.csv")
    args = parser.parse_args()

    ft_articles = collect_ft_articles(args.ft_dir)
    available_pos = collect_available_pos(args.po_dir)
    far_rows = read_far_rows(args.far_list)
    usable, stats = build_usable_rows(far_rows, ft_articles, available_pos)
    write_csv(usable, args.out)

    print(f"FT articles on hand:        {len(ft_articles)}")
    print(f"PO-linked FT files on hand: {len(available_pos)} (excluding {sorted(EXCLUDED_POS)})")
    print(f"Total FAR rows:             {stats['total_far_rows']}")
    print(f"  - dropped, no deviation:  {stats['dropped_no_deviation']}")
    print(f"  - dropped, no FT article: {stats['dropped_no_ft_article']}")
    print(f"  - dropped, excluded PO:   {stats['dropped_excluded_po']}")
    print(f"Usable rows:                {stats['usable_rows']}")
    print(f"  - PO already covered:     {stats['po_covered']}")
    print(f"Written to {args.out}")


if __name__ == "__main__":
    main()
